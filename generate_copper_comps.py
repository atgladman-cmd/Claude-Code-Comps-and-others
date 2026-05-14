#!/usr/bin/env python3
"""
ASX Copper & Base Metals Comparable Companies Analysis
Generates copper_comps_table.xlsx with live IRESS RTD formulas.

Sheets:
  1. Comps Table   – main institutional-quality comps table
  2. Assumptions   – metal prices, CuEq calculations, data sources
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ──────────────────────────────────────────────────────────────────────────────
# COLOUR PALETTE
# ──────────────────────────────────────────────────────────────────────────────
C_NAVY       = "1F4E79"   # deep navy  – title / primary headers
C_BLUE       = "2E75B6"   # mid-blue   – sub-headers
C_LTBLUE     = "D6E4F7"   # pale blue  – odd data row band
C_WHITE      = "FFFFFF"
C_AMBER      = "FFF2CC"   # amber      – footnote / flag highlight
C_AMBER_DARK = "7F6000"   # dark amber – footnote text
C_LGREY      = "F2F2F2"   # light grey – source section bg
C_DGREY      = "595959"
C_BLACK      = "1A1A1A"
C_BORDER     = "9DC3E6"   # light blue border
C_NAVY2      = "17375E"   # assumptions section header
C_BGASMP     = "EBF3FB"   # assumptions section bg

# Number format strings
FMT_PRICE  = '"A$"#,##0.000'
FMT_AUD_M  = '"A$"#,##0.0'
FMT_AUD_T  = '"A$"#,##0'
FMT_KT     = "#,##0.0"
FMT_SHS    = "#,##0.0"
FMT_ACCT   = '_-"A$"* #,##0.0_-;[Red]_-"A$"* (#,##0.0)_-;_-"A$"* "-"??_-;_-@_-'
FMT_ACCT_T = '_-"A$"* #,##0_-;[Red]_-"A$"* (#,##0)_-;_-"A$"* "-"??_-;_-@_-'
FMT_USD    = '"US$"#,##0.0'
FMT_4DP    = "#,##0.0000"


# ──────────────────────────────────────────────────────────────────────────────
# STYLE HELPERS
# ──────────────────────────────────────────────────────────────────────────────
def fnt(bold=False, sz=10, col=C_BLACK, italic=False):
    return Font(name="Calibri", bold=bold, size=sz, color=col, italic=italic)

def pfill(hex_col):
    return PatternFill("solid", fgColor=hex_col)

def aln(h="left", v="center", wrap=False, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)

def bdr(style="thin", col=C_BORDER):
    s = Side(style=style, color=col)
    return Border(left=s, right=s, top=s, bottom=s)

def bdr_thick_bot():
    t = Side(style="thin",   color=C_BORDER)
    b = Side(style="medium", color=C_NAVY)
    return Border(left=t, right=t, top=t, bottom=b)

def bdr_none():
    return Border()

def style(cell, bold=False, sz=10, col=C_BLACK, italic=False,
          bg=None, h="left", v="center", wrap=False,
          num_fmt=None, border_style="thin", thick_bot=False, no_border=False):
    cell.font      = fnt(bold=bold, sz=sz, col=col, italic=italic)
    cell.alignment = aln(h=h, v=v, wrap=wrap)
    if bg:
        cell.fill = pfill(bg)
    if no_border:
        cell.border = bdr_none()
    elif thick_bot:
        cell.border = bdr_thick_bot()
    else:
        cell.border = bdr(style=border_style, col=C_BORDER)
    if num_fmt:
        cell.number_format = num_fmt


def merge_title(ws, row, col_start, col_end, value,
                bold=True, sz=13, fg=C_WHITE, bg=C_NAVY,
                h="center", v="center", row_ht=None):
    ws.merge_cells(start_row=row, start_column=col_start,
                   end_row=row,   end_column=col_end)
    c = ws.cell(row=row, column=col_start, value=value)
    style(c, bold=bold, sz=sz, col=fg, bg=bg, h=h, v=v, no_border=True)
    if row_ht:
        ws.row_dimensions[row].height = row_ht
    return c


def write(ws, row, col, val=None, formula=None,
          bold=False, sz=10, col_txt=C_BLACK, italic=False,
          bg=None, h="left", v="center", wrap=False,
          num_fmt=None, thick_bot=False, no_border=False):
    cell = ws.cell(row=row, column=col)
    if formula is not None:
        cell.value = formula
    elif val is not None:
        cell.value = val
    style(cell, bold=bold, sz=sz, col=col_txt, italic=italic,
          bg=bg, h=h, v=v, wrap=wrap, num_fmt=num_fmt,
          thick_bot=thick_bot, no_border=no_border)
    return cell


# ──────────────────────────────────────────────────────────────────────────────
# COMPANY MASTER DATA
# ──────────────────────────────────────────────────────────────────────────────
COMPANIES = [
    {
        "ticker":    "AIS",
        "name":      "Aeris Resources Limited",
        "shares_m":  1150.0,
        "shares_src":"Investor Pres / Q3 FY26 Quarterly Activities Report, Apr-2026",
        "cash_m":    149.8,
        "debt_m":    0.0,
        "qtr":       "31-Mar-26",
        # Individual metal guidance midpoints
        "cu_kt":     26.5,    # Tritton: 24–29 kt Cu
        "au_koz":    41.0,    # Cracow: 36–46 koz Au
        "ag_koz":    0.0,
        "zn_kt":     0.0,
        "pb_kt":     0.0,
        # Company-disclosed CuEq midpoint (0 = not disclosed; calc will be used)
        "disc_cueq": 44.5,    # 40–49 kt CuEq FY2026
        "flag":      "†",     # † = company-disclosed CuEq
        "gsrc":      "FY2026 production guidance (40–49 kt CuEq); Q3 FY26 Quarterly Activities Report, Apr-2026",
        "note":      "Debt-free from Dec-2025. Tritton Cu mine (NSW) + Cracow Au mine (Qld).",
    },
    {
        "ticker":    "A1M",
        "name":      "AIC Mines Limited",
        "shares_m":  798.0,
        "shares_src":"Q3 FY26 Quarterly Report (797.6M at 28-Mar-26), Apr-2026",
        "cash_m":    31.1,
        "debt_m":    45.3,
        "qtr":       "31-Mar-26",
        "cu_kt":     13.0,    # 12,800–13,200 t Cu
        "au_koz":    6.25,    # 6,000–6,500 oz Au (Eloise co-product)
        "ag_koz":    0.0,
        "zn_kt":     0.0,
        "pb_kt":     0.0,
        "disc_cueq": 0.0,
        "flag":      "‡",     # ‡ = CuEq calculated using live spot prices
        "gsrc":      "FY2026 guidance (12,800–13,200 t Cu + 6,000–6,500 oz Au); Q3 FY26 Quarterly, Apr-2026",
        "note":      "Eloise Cu mine (NW Qld). Q3 cash depressed – A$8.3M concentrate on-site, shipped Apr. Trafigura facility A$45.3M.",
    },
    {
        "ticker":    "DVP",
        "name":      "Develop Global Limited",
        "shares_m":  330.0,
        "shares_src":"Q3 FY26 Quarterly / market data, Apr-2026 (~329.99M shares)",
        "cash_m":    180.0,
        "debt_m":    158.0,
        "qtr":       "31-Mar-26",
        "cu_kt":     10.0,    # Woodlawn steady-state ~10 ktpa Cu in concentrate
        "au_koz":    0.0,
        "ag_koz":    0.0,
        "zn_kt":     30.0,    # Woodlawn ~30+ ktpa Zn in concentrate
        "pb_kt":     0.0,
        "disc_cueq": 0.0,
        "flag":      "‡",
        "gsrc":      "Woodlawn LOM run-rate (~10 ktpa Cu, ~30 ktpa Zn); FY2026 specific annual guidance not formally published; Q3 FY26 Quarterly, Apr-2026",
        "note":      "Woodlawn (NSW): Cu/Zn/Pb/Ag polymetallic. Also has large mining services division (>A$50M/qtr revenue). Pb & Ag guidance not quantified.",
    },
    {
        "ticker":    "29M",
        "name":      "29Metals Limited",
        "shares_m":  1747.3,
        "shares_src":"Post-entitlement offer (374.95M new shares Feb-2026); Q1 CY26 Quarterly, Apr-2026",
        "cash_m":    234.0,   # Adjusted to reconcile to company-disclosed net cash A$48M (31-Mar-26)
        "debt_m":    186.0,   # US$120M RCF ≈ A$186M @ 0.645 AUDUSD
        "qtr":       "31-Mar-26",
        "cu_kt":     22.0,    # 20–24 kt Cu
        "au_koz":    10.0,    # 6–14 koz Au (midpoint)
        "ag_koz":    500.0,   # 400–600 koz Ag (midpoint)
        "zn_kt":     15.0,    # 5–25 kt Zn (midpoint; reduced from 40–50 kt after Xantho disruption)
        "pb_kt":     0.0,
        "disc_cueq": 0.0,
        "flag":      "‡",
        "gsrc":      "Revised FY2026 guidance; Q1 CY26 Quarterly, Apr-2026 (Cu 20–24 kt; Zn 5–25 kt; Au 6–14 koz; Ag 400–600 koz)",
        "note":      "Golden Grove (WA): Cu–Zn–Au–Ag. Zn guidance materially cut (Xantho Extended geotechnical event). Net cash A$48M per co. (31-Mar-26). US$120M RCF converted @ 0.645.",
    },
    {
        "ticker":    "AMI",
        "name":      "Aurelia Metals Limited",
        "shares_m":  1693.0,
        "shares_src":"Q3 FY26 Quarterly / market data, Apr-2026",
        "cash_m":    94.7,
        "debt_m":    8.6,
        "qtr":       "31-Mar-26",
        "cu_kt":     2.75,    # 2.5–3.0 kt Cu (minor co-product, midpoint)
        "au_koz":    47.5,    # 45–50 koz Au (upgraded guidance, primary metal)
        "ag_koz":    0.0,
        "zn_kt":     28.0,    # 24–32 kt Zn
        "pb_kt":     18.0,    # 14–22 kt Pb
        "disc_cueq": 0.0,
        "flag":      "‡",
        "gsrc":      "Upgraded FY2026 guidance; Q3 FY26 Quarterly, Apr-2026 (Cu 2.5–3.0 kt; Au 45–50 koz; Zn 24–32 kt; Pb 14–22 kt)",
        "note":      "Federation + Peak mines (Cobar, NSW). Au is primary revenue driver; Cu is minor co-product. A$150M facility committed but undrawn at 31-Mar-26.",
    },
    {
        "ticker":    "HGO",
        "name":      "Hillgrove Resources Limited",
        "shares_m":  3410.0,
        "shares_src":"Q1 CY26 Quarterly / market data, Apr-2026",
        "cash_m":    25.2,
        "debt_m":    0.0,
        "qtr":       "31-Mar-26",
        "cu_kt":     13.375,  # CY2026 guidance: 12,750–14,000 t Cu (midpoint)
        "au_koz":    0.0,
        "ag_koz":    0.0,
        "zn_kt":     0.0,
        "pb_kt":     0.0,
        "disc_cueq": 0.0,
        "flag":      "",      # No flag – pure copper, no conversion required
        "gsrc":      "CY2026 guidance (12,750–14,000 t Cu); Q4 CY25 Quarterly Activities Report, Jan-2026",
        "note":      "Kanmantoo underground (SA). Pure copper producer (cathode + concentrate). Debt-free. Record Q1 CY26: 3,120 t Cu.",
    },
    {
        "ticker":    "CYM",
        "name":      "Cyprium Metals Limited",
        "shares_m":  492.2,
        "shares_src":"Q1 CY26 Quarterly; post A$80M placement + A$41M further raise; post 1:10 consolidation Oct-2025; Apr-2026",
        "cash_m":    96.5,
        "debt_m":    42.0,    # US$27.3M Nebari facility ≈ A$42M @ 0.645
        "qtr":       "31-Mar-26",
        "cu_kt":     0.0,
        "au_koz":    0.0,
        "ag_koz":    0.0,
        "zn_kt":     0.0,
        "pb_kt":     0.0,
        "disc_cueq": 7.0,     # Phase 1 run-rate ~7 ktpa Cu cathode (annualised)
        "flag":      "§",     # § = pre-production; annualised Phase 1 run-rate
        "gsrc":      "Phase 1 Nifty Copper Complex restart targeting Q3 CY26; run-rate ~7 ktpa Cu cathode; Q1 CY26 Quarterly, Apr-2026",
        "note":      "PRE-PRODUCTION. Nifty SX-EW cathode restart (Pilbara, WA). Phase 1 ~7 ktpa Cu; Phase 2 potential ~25 ktpa. US$27.3M Nebari facility ≈ A$42M @ 0.645.",
    },
    {
        "ticker":    "CSC",
        "name":      "Capstone Copper Corp.",
        "shares_m":  763.7,
        "shares_src":"Q1 CY26 MD&A (Canadian disclosure), 29-Apr-2026",
        "cash_m":    611.0,   # US$394.1M converted @ 0.645 AUDUSD
        "debt_m":    1754.0,  # US$1,131.8M converted @ 0.645 AUDUSD
        "qtr":       "31-Mar-26",
        "cu_kt":     215.0,   # CY2026: 200–230 kt Cu (midpoint; copper-only figure)
        "au_koz":    0.0,
        "ag_koz":    0.0,
        "zn_kt":     0.0,
        "pb_kt":     0.0,
        "disc_cueq": 0.0,
        "flag":      "‡",
        "gsrc":      "CY2026 production guidance (200–230 kt Cu); Q1 CY26 Results, 29-Apr-2026 (Pinto Valley AZ + Cozamin MX + Mantos Blancos CL + Mantoverde CL)",
        "note":      "Canadian co. / ASX CDI (1:1). All financials in USD, converted @ 0.645 AUDUSD. Files Canadian MD&A (not App 5B). Co-products (Ag, Au, Mo) add incremental CuEq – not quantified here.",
    },
    {
        "ticker":    "SFR",
        "name":      "Sandfire Resources Limited",
        "shares_m":  466.6,
        "shares_src":"Q3 FY26 Quarterly (466.58M at 28-Mar-26), 23-Apr-2026",
        "cash_m":    195.3,   # Implied: A$76M net cash + A$119.3M debt
        "debt_m":    119.3,
        "qtr":       "31-Mar-26",
        "cu_kt":     0.0,
        "au_koz":    0.0,
        "ag_koz":    0.0,
        "zn_kt":     0.0,
        "pb_kt":     0.0,
        "disc_cueq": 157.0,   # 149–165 kt CuEq FY2026 (midpoint; company-disclosed)
        "flag":      "†",
        "gsrc":      "FY2026 guidance (149–165 kt CuEq; guiding to lower half); Q3 FY26 Quarterly, 23-Apr-2026 (MATSA Spain + Motheo Botswana)",
        "note":      "Company-disclosed CuEq. MATSA: Cu/Zn/Pb/Ag polymetallic. Motheo: copper only. Net cash A$76M per co. (31-Mar-26).",
    },
]

# Row anchors on the Assumptions sheet (1-based)
ASMPT_TITLE_ROW     = 1
ASMPT_PRICE_HDR     = 4
ASMPT_CU_ROW        = 5
ASMPT_AU_ROW        = 6
ASMPT_AG_ROW        = 7
ASMPT_ZN_ROW        = 8
ASMPT_PB_ROW        = 9
ASMPT_FX_ROW        = 10
ASMPT_CUEQ_HDR      = 14   # "CuEq Calculations" section header
ASMPT_CUEQ_COL_HDR  = 15   # column headers for CuEq table
ASMPT_CUEQ_DATA     = 16   # first company row in CuEq table
ASMPT_SRC_HDR       = 28   # Data Sources section
ASMPT_SRC_SHS_HDR   = 30
ASMPT_SRC_SHS_DATA  = 31
ASMPT_SRC_ND_HDR    = 43
ASMPT_SRC_ND_DATA   = 44

# CuEq table columns (1-based, on Assumptions sheet)
C_TICKER   = 1   # A
C_CU_KT    = 2   # B  Cu guidance midpoint (kt)
C_AU_KOZ   = 3   # C  Au guidance midpoint (koz)
C_AG_KOZ   = 4   # D  Ag guidance midpoint (koz)
C_ZN_KT    = 5   # E  Zn guidance midpoint (kt)
C_PB_KT    = 6   # F  Pb guidance midpoint (kt)
C_CU_CUEQ  = 7   # G  Cu CuEq contribution (kt) = Cu guidance
C_AU_CUEQ  = 8   # H  Au CuEq contribution (kt) = C * Au_price / Cu_price
C_AG_CUEQ  = 9   # I  Ag CuEq contribution (kt) = D * Ag_price / Cu_price
C_ZN_CUEQ  = 10  # J  Zn CuEq contribution (kt) = E * Zn_price / Cu_price
C_PB_CUEQ  = 11  # K  Pb CuEq contribution (kt) = F * Pb_price / Cu_price
C_CALC_TOT = 12  # L  Calculated total CuEq (kt) = SUM(G:K)
C_DISC_CUEQ= 13  # M  Company-disclosed CuEq (kt) [0 if not disclosed]
C_CUEQ_USED= 14  # N  =IF(M>0,M,L)  ← referenced by Comps sheet
C_FLAG_COL = 15  # O  Flag character
C_GSRC_COL = 16  # P  Guidance source text


# ──────────────────────────────────────────────────────────────────────────────
# BUILD ASSUMPTIONS SHEET
# ──────────────────────────────────────────────────────────────────────────────
def build_assumptions(ws):
    # Column widths
    col_widths = {1: 8, 2: 16, 3: 22, 4: 12, 5: 12, 6: 12,
                  7: 12, 8: 12, 9: 12, 10: 12, 11: 12,
                  12: 13, 13: 13, 14: 13, 15: 6, 16: 60}
    for col, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    # ── Title ─────────────────────────────────────────────────────────────────
    merge_title(ws, ASMPT_TITLE_ROW, 1, 16,
                "ASSUMPTIONS, CuEq CALCULATIONS & DATA SOURCES",
                sz=12, row_ht=22)
    merge_title(ws, 2, 1, 16,
                "ASX Copper & Base Metals Comps | Data sourced from ASX company announcements (asx.com.au)",
                bold=False, sz=9, bg=C_BLUE, row_ht=14)

    # ── Section 1: Metal Price Assumptions ────────────────────────────────────
    merge_title(ws, 3, 1, 16, "  SECTION 1 — METAL PRICE ASSUMPTIONS  (live via IRESS RTD)",
                sz=10, bg=C_NAVY2, row_ht=16)

    # Header row
    hdrs = ["Metal", "IRESS Symbol", "Live Price", "Unit", "Notes / Description"]
    for i, h in enumerate(hdrs, 1):
        write(ws, ASMPT_PRICE_HDR, i, val=h,
              bold=True, col_txt=C_WHITE, bg=C_BLUE, h="center", sz=9, thick_bot=True)

    # Price rows
    prices = [
        ("Copper",  "COPNY.ID",  '=@IRESSRtd("Quote","13","0","COPNY.ID")',  "USD/t",   "LME Copper (Grade A)"),
        ("Gold",    "SPTGLD.IF", '=@IRESSRtd("Quote","13","0","SPTGLD.IF")', "USD/oz",  "Spot Gold (LBMA PM fix)"),
        ("Silver",  "SPTSLV.IF", '=@IRESSRtd("Quote","13","0","SPTSLV.IF")', "USD/oz",  "Spot Silver"),
        ("Zinc",    "ZINC.LME",  '=@IRESSRtd("Quote","13","0","ZINC.LME")',  "USD/t",   "LME Cash Zinc"),
        ("Lead",    "LEAD.LME",  '=@IRESSRtd("Quote","13","0","LEAD.LME")',  "USD/t",   "LME Cash Lead"),
        ("AUD/USD", "AUDUSD.FX", '=@IRESSRtd("Quote","13","0","AUDUSD.FX")', "AUD/USD", "Spot FX – used for USD→AUD conversion"),
    ]
    price_fmts = ['"US$"#,##0.00', '"US$"#,##0.00', '"US$"#,##0.0000',
                  '"US$"#,##0.00', '"US$"#,##0.00', "0.0000"]
    for idx, ((metal, sym, rtd, unit, note), pfmt) in enumerate(zip(prices, price_fmts)):
        r = ASMPT_CU_ROW + idx
        bg = C_BGASMP if idx % 2 == 0 else C_WHITE
        write(ws, r, 1, val=metal,  bg=bg, sz=9, bold=True)
        write(ws, r, 2, val=sym,    bg=bg, sz=9, h="center")
        write(ws, r, 3, formula=rtd, bg=bg, sz=10, bold=True, h="right", num_fmt=pfmt)
        write(ws, r, 4, val=unit,   bg=bg, sz=9, h="center")
        write(ws, r, 5, val=note,   bg=bg, sz=9)

    # ── Section 2: CuEq Calculations ─────────────────────────────────────────
    ws.row_dimensions[12].height = 6
    merge_title(ws, 13, 1, 16, "  SECTION 2 — CuEq PRODUCTION CALCULATIONS",
                sz=10, bg=C_NAVY2, row_ht=16)

    note_txt = (
        "CuEq methodology: CuEq (kt) = Cu_kt + Σ [ By-product_guidance × by-product_price / copper_price ]  |  "
        "Gold & Silver in koz (×1,000 oz/koz); Cu, Zn, Pb prices in USD/t; Au, Ag prices in USD/oz  |  "
        "Units resolve to kt: (koz × USD/oz) / (USD/t) = kt  |  "
        "† Company-disclosed CuEq used (col M).  ‡ Calculated from components (col L).  § Pre-production annualised run-rate."
    )
    merge_title(ws, ASMPT_CUEQ_HDR, 1, 16, note_txt,
                bold=False, sz=8, bg=C_AMBER, fg=C_AMBER_DARK,
                h="left", row_ht=30)
    ws.cell(row=ASMPT_CUEQ_HDR, column=1).alignment = aln(h="left", v="center", wrap=True)

    # CuEq column headers
    cueq_hdrs = [
        "Ticker", "Cu Guide\n(kt mid)", "Au Guide\n(koz mid)",
        "Ag Guide\n(koz mid)", "Zn Guide\n(kt mid)", "Pb Guide\n(kt mid)",
        "Cu→CuEq\n(kt)", "Au→CuEq\n(kt)", "Ag→CuEq\n(kt)",
        "Zn→CuEq\n(kt)", "Pb→CuEq\n(kt)",
        "CALC\nCuEq (kt)", "DISC\nCuEq (kt)", "CuEq USED\n(kt)",
        "Flag", "Guidance Source"
    ]
    ws.row_dimensions[ASMPT_CUEQ_COL_HDR].height = 30
    for ci, h in enumerate(cueq_hdrs, 1):
        write(ws, ASMPT_CUEQ_COL_HDR, ci, val=h,
              bold=True, col_txt=C_WHITE, bg=C_BLUE, h="center",
              v="center", wrap=True, sz=8, thick_bot=True)

    # Price cell absolute references (on Assumptions sheet)
    cu_ref = f"$C${ASMPT_CU_ROW}"
    au_ref = f"$C${ASMPT_AU_ROW}"
    ag_ref = f"$C${ASMPT_AG_ROW}"
    zn_ref = f"$C${ASMPT_ZN_ROW}"
    pb_ref = f"$C${ASMPT_PB_ROW}"

    for i, co in enumerate(COMPANIES):
        r = ASMPT_CUEQ_DATA + i
        bg = C_BGASMP if i % 2 == 0 else C_WHITE

        # Column formulas
        bc = get_column_letter(C_CU_KT)
        cc = get_column_letter(C_AU_KOZ)
        dc = get_column_letter(C_AG_KOZ)
        ec = get_column_letter(C_ZN_KT)
        fc = get_column_letter(C_PB_KT)
        gc = get_column_letter(C_CU_CUEQ)
        hc = get_column_letter(C_AU_CUEQ)
        ic = get_column_letter(C_AG_CUEQ)
        jc = get_column_letter(C_ZN_CUEQ)
        kc = get_column_letter(C_PB_CUEQ)
        lc = get_column_letter(C_CALC_TOT)
        mc = get_column_letter(C_DISC_CUEQ)

        write(ws, r, C_TICKER,    val=co["ticker"],     bg=bg, sz=9, bold=True)
        write(ws, r, C_CU_KT,     val=co["cu_kt"],      bg=bg, sz=9, h="right", num_fmt=FMT_KT)
        write(ws, r, C_AU_KOZ,    val=co["au_koz"],     bg=bg, sz=9, h="right", num_fmt=FMT_KT)
        write(ws, r, C_AG_KOZ,    val=co["ag_koz"],     bg=bg, sz=9, h="right", num_fmt=FMT_KT)
        write(ws, r, C_ZN_KT,     val=co["zn_kt"],      bg=bg, sz=9, h="right", num_fmt=FMT_KT)
        write(ws, r, C_PB_KT,     val=co["pb_kt"],      bg=bg, sz=9, h="right", num_fmt=FMT_KT)
        # Calculated CuEq contributions
        write(ws, r, C_CU_CUEQ,   formula=f"={bc}{r}",              bg=bg, sz=9, h="right", num_fmt=FMT_KT)
        write(ws, r, C_AU_CUEQ,   formula=f"={cc}{r}*{au_ref}/{cu_ref}", bg=bg, sz=9, h="right", num_fmt=FMT_KT)
        write(ws, r, C_AG_CUEQ,   formula=f"={dc}{r}*{ag_ref}/{cu_ref}", bg=bg, sz=9, h="right", num_fmt=FMT_KT)
        write(ws, r, C_ZN_CUEQ,   formula=f"={ec}{r}*{zn_ref}/{cu_ref}", bg=bg, sz=9, h="right", num_fmt=FMT_KT)
        write(ws, r, C_PB_CUEQ,   formula=f"={fc}{r}*{pb_ref}/{cu_ref}", bg=bg, sz=9, h="right", num_fmt=FMT_KT)
        write(ws, r, C_CALC_TOT,  formula=f"=SUM({gc}{r}:{kc}{r})",      bg=bg, sz=9, bold=True, h="right", num_fmt=FMT_KT)
        write(ws, r, C_DISC_CUEQ, val=co["disc_cueq"] if co["disc_cueq"] > 0 else None,
              bg=bg, sz=9, h="right", num_fmt=FMT_KT)
        write(ws, r, C_CUEQ_USED, formula=f"=IF({mc}{r}>0,{mc}{r},{lc}{r})",
              bg=C_LTBLUE, sz=9, bold=True, h="right", num_fmt=FMT_KT)
        write(ws, r, C_FLAG_COL,  val=co["flag"],        bg=bg, sz=9, h="center")
        write(ws, r, C_GSRC_COL,  val=co["gsrc"],        bg=bg, sz=8, italic=True, wrap=True)
        ws.row_dimensions[r].height = 28

    # ── Section 3: Data Sources ───────────────────────────────────────────────
    ws.row_dimensions[ASMPT_SRC_HDR - 1].height = 6
    merge_title(ws, ASMPT_SRC_HDR, 1, 16,
                "  SECTION 3 — DATA SOURCES",
                sz=10, bg=C_NAVY2, row_ht=16)
    merge_title(ws, ASMPT_SRC_HDR + 1, 1, 16,
                "All data sourced from ASX company announcements (asx.com.au) and company investor relations pages.  "
                "All figures in A$M unless noted.  USD-denominated items converted at 0.645 AUDUSD (implied rate at sourcing).",
                bold=False, sz=8, bg=C_AMBER, fg=C_AMBER_DARK, h="left", row_ht=20)
    ws.cell(row=ASMPT_SRC_HDR+1, column=1).alignment = aln(h="left", v="center", wrap=True)

    # ── 3a: Shares on issue ───────────────────────────────────────────────────
    merge_title(ws, ASMPT_SRC_SHS_HDR, 1, 16,
                "3A. SHARES ON ISSUE", bold=True, sz=9, bg=C_NAVY, row_ht=14)
    shs_hdrs = ["Ticker", "Company", "Shares (M)", "Source Document", "Date", "Doc Type"]
    for ci, h in enumerate(shs_hdrs, 1):
        write(ws, ASMPT_SRC_SHS_HDR + 1, ci, val=h,
              bold=True, col_txt=C_WHITE, bg=C_BLUE, h="center", sz=9)
    ws.merge_cells(start_row=ASMPT_SRC_SHS_HDR+1, start_column=4,
                   end_row=ASMPT_SRC_SHS_HDR+1, end_column=6)

    for i, co in enumerate(COMPANIES):
        r = ASMPT_SRC_SHS_DATA + i
        bg = C_LGREY if i % 2 == 0 else C_WHITE
        write(ws, r, 1, val=co["ticker"],    bg=bg, sz=9, bold=True)
        write(ws, r, 2, val=co["name"],      bg=bg, sz=9)
        write(ws, r, 3, val=co["shares_m"],  bg=bg, sz=9, h="right", num_fmt=FMT_SHS)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
        write(ws, r, 4, val=co["shares_src"], bg=bg, sz=8, italic=True, wrap=True)
        ws.row_dimensions[r].height = 18

    # ── 3b: Net Cash / Debt ───────────────────────────────────────────────────
    r_nd_hdr = ASMPT_SRC_SHS_DATA + len(COMPANIES) + 2
    merge_title(ws, r_nd_hdr, 1, 16,
                "3B. NET CASH / (DEBT) — MOST RECENT QUARTERLY (APP 5B OR EQUIVALENT)",
                bold=True, sz=9, bg=C_NAVY, row_ht=14)
    nd_hdrs = ["Ticker", "Company", "Cash (A$M)", "Debt (A$M)", "Net (Cash)/Debt (A$M)", "Qtr End", "Source"]
    for ci, h in enumerate(nd_hdrs, 1):
        write(ws, r_nd_hdr + 1, ci, val=h,
              bold=True, col_txt=C_WHITE, bg=C_BLUE, h="center", sz=9)

    for i, co in enumerate(COMPANIES):
        r = r_nd_hdr + 2 + i
        net_debt = co["debt_m"] - co["cash_m"]
        bg = C_LGREY if i % 2 == 0 else C_WHITE
        write(ws, r, 1, val=co["ticker"],      bg=bg, sz=9, bold=True)
        write(ws, r, 2, val=co["name"],        bg=bg, sz=9)
        write(ws, r, 3, val=co["cash_m"],      bg=bg, sz=9, h="right", num_fmt=FMT_AUD_M)
        write(ws, r, 4, val=co["debt_m"],      bg=bg, sz=9, h="right", num_fmt=FMT_AUD_M)
        write(ws, r, 5, val=net_debt,          bg=bg, sz=9, h="right", num_fmt=FMT_ACCT,
              bold=True)
        write(ws, r, 6, val=co["qtr"],         bg=bg, sz=9, h="center")
        write(ws, r, 7, val=co["note"],        bg=bg, sz=8, italic=True, wrap=True)
        ws.row_dimensions[r].height = 22

    # Freeze top rows
    ws.freeze_panes = ws.cell(row=ASMPT_PRICE_HDR + 1, column=1)


# ──────────────────────────────────────────────────────────────────────────────
# BUILD MAIN COMPS SHEET
# ──────────────────────────────────────────────────────────────────────────────
COMPS_TITLE_ROW  = 1
COMPS_SUBTITLE   = 2
COMPS_HDR_ROW    = 4
COMPS_DATA_START = 5   # first company row
NUM_COMPANIES    = len(COMPANIES)

# Comps sheet columns
CC_NUM    = 1    # A  Row number
CC_NAME   = 2    # B  Company name
CC_TICKER = 3    # C  ASX code
CC_PRICE  = 4    # D  Share price (IRESS RTD)
CC_SHS    = 5    # E  Shares on issue (M)
CC_MKTCAP = 6    # F  Market cap (A$M)
CC_NETDEBT= 7    # G  Net (cash)/debt (A$M)
CC_EV     = 8    # H  EV (A$M)
CC_CUEQ   = 9    # I  CuEq guidance (kt pa)
CC_EVCUEQ = 10   # J  EV / CuEq (A$/t)
CC_FLAGS  = 11   # K  Flags / notes col

TOTAL_COLS = 11


def build_comps(ws, ws_asmpt_name="Assumptions"):
    # ── Column widths ─────────────────────────────────────────────────────────
    col_widths = {
        CC_NUM:     4,
        CC_NAME:    29,
        CC_TICKER:  7,
        CC_PRICE:   13,
        CC_SHS:     13,
        CC_MKTCAP:  13,
        CC_NETDEBT: 16,
        CC_EV:      13,
        CC_CUEQ:    14,
        CC_EVCUEQ:  13,
        CC_FLAGS:   5,
    }
    for col, w in col_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    # ── Title row ─────────────────────────────────────────────────────────────
    ws.row_dimensions[COMPS_TITLE_ROW].height = 26
    merge_title(ws, COMPS_TITLE_ROW, 1, TOTAL_COLS,
                "ASX COPPER & BASE METALS — COMPARABLE COMPANIES ANALYSIS",
                sz=14, row_ht=26)

    ws.row_dimensions[COMPS_SUBTITLE].height = 14
    merge_title(ws, COMPS_SUBTITLE, 1, TOTAL_COLS,
                "Share prices & metal spot prices sourced live via IRESS RTD  |  "
                "Financial data from ASX company announcements (Appendix 5B / Quarterly Activities Reports) — March 2026 quarter  |  "
                "All figures in A$M unless stated  |  Guidance midpoints used throughout",
                bold=False, sz=9, bg=C_BLUE, row_ht=14)
    ws.cell(row=COMPS_SUBTITLE, column=1).alignment = aln(h="left", v="center", wrap=False)

    # ── Column headers ────────────────────────────────────────────────────────
    ws.row_dimensions[3].height = 4   # spacer
    ws.row_dimensions[COMPS_HDR_ROW].height = 36
    headers = [
        ("#",             "center"),
        ("Company Name",  "left"),
        ("ASX\nCode",     "center"),
        ("Share Price\n(A$/share)",  "right"),
        ("Shares on\nIssue (M)",     "right"),
        ("Market Cap\n(A$M)",        "right"),
        ("Net (Cash) /\nDebt (A$M)", "right"),
        ("EV\n(A$M)",                "right"),
        ("CuEq Guidance\n(kt pa) ⁽¹⁾", "right"),
        ("EV / CuEq\n(A$/t)",        "right"),
        ("⚑",            "center"),
    ]
    for ci, (hdr, halign) in enumerate(headers, 1):
        c = write(ws, COMPS_HDR_ROW, ci, val=hdr,
                  bold=True, col_txt=C_WHITE, bg=C_NAVY,
                  h=halign, v="center", wrap=True, sz=9, thick_bot=True)

    # ── Data rows ─────────────────────────────────────────────────────────────
    for i, co in enumerate(COMPANIES):
        r = COMPS_DATA_START + i
        bg = C_LTBLUE if i % 2 == 0 else C_WHITE
        ws.row_dimensions[r].height = 18

        asmpt_row = ASMPT_CUEQ_DATA + i
        cueq_ref  = f"=Assumptions!{get_column_letter(C_CUEQ_USED)}{asmpt_row}"
        net_debt  = round(co["debt_m"] - co["cash_m"], 1)

        # Row number
        write(ws, r, CC_NUM, val=i + 1,
              h="center", bg=bg, sz=9, col_txt=C_DGREY)

        # Company name
        write(ws, r, CC_NAME, val=co["name"],
              bold=(i < 2), bg=bg, sz=10)

        # Ticker
        write(ws, r, CC_TICKER, val=co["ticker"],
              bold=True, h="center", bg=bg, sz=10)

        # Share price – live IRESS RTD
        write(ws, r, CC_PRICE,
              formula=f'=@IRESSRtd("Quote","13","0","{co["ticker"]}")',
              h="right", bg=bg, sz=10, bold=True, num_fmt=FMT_PRICE)

        # Shares on issue (M)
        write(ws, r, CC_SHS, val=co["shares_m"],
              h="right", bg=bg, sz=10, num_fmt=FMT_SHS)

        # Market cap = Price × Shares (in A$M: A$/sh × Mshares = A$M directly)
        dc = get_column_letter(CC_PRICE)
        ec = get_column_letter(CC_SHS)
        write(ws, r, CC_MKTCAP,
              formula=f"={dc}{r}*{ec}{r}",
              h="right", bg=bg, sz=10, num_fmt=FMT_AUD_M)

        # Net (cash) / debt  [positive = debt, negative = cash]
        write(ws, r, CC_NETDEBT, val=net_debt,
              h="right", bg=bg, sz=10, num_fmt=FMT_ACCT)

        # EV = Mkt Cap + Net Debt
        fc = get_column_letter(CC_MKTCAP)
        gc = get_column_letter(CC_NETDEBT)
        write(ws, r, CC_EV,
              formula=f"={fc}{r}+{gc}{r}",
              h="right", bg=bg, sz=10, bold=True, num_fmt=FMT_AUD_M)

        # CuEq production guidance (kt) – live from Assumptions sheet
        write(ws, r, CC_CUEQ,
              formula=cueq_ref,
              h="right", bg=bg, sz=10, num_fmt=FMT_KT)

        # EV / CuEq (A$/t) = EV(A$M) × 1,000,000 / (CuEq(kt) × 1,000)
        #                   = EV(A$M) × 1,000 / CuEq(kt)
        hc = get_column_letter(CC_EV)
        ic = get_column_letter(CC_CUEQ)
        write(ws, r, CC_EVCUEQ,
              formula=f'=IF({ic}{r}>0,{hc}{r}*1000/{ic}{r},"n/a")',
              h="right", bg=bg, sz=10, bold=True, num_fmt=FMT_AUD_T)

        # Flag column
        write(ws, r, CC_FLAGS, val=co["flag"],
              h="center", bg=bg, sz=9,
              col_txt=C_AMBER_DARK if co["flag"] else C_DGREY)

    # ── Footnotes section ─────────────────────────────────────────────────────
    fn_row = COMPS_DATA_START + NUM_COMPANIES + 1

    # Thin divider
    for col in range(1, TOTAL_COLS + 1):
        c = ws.cell(row=fn_row - 1, column=col)
        c.border = Border(bottom=Side(style="medium", color=C_NAVY))

    ws.row_dimensions[fn_row].height = 4  # spacer

    footnotes = [
        ("⁽¹⁾  CuEq Guidance",
         "Guidance midpoints used. Polymetallic CuEq calculated using live IRESS spot prices (see Assumptions sheet, Section 2). "
         "Guidance ranges and data sources are detailed in the Assumptions sheet."),
        ("†  Company-disclosed CuEq",
         "AIS (40–49 kt, FY2026) and SFR (149–165 kt, FY2026): CuEq as disclosed by the company in quarterly activities reports. "
         "Each company's own price assumptions are embedded in those figures."),
        ("‡  Calculated CuEq",
         "A1M, DVP, 29M, AMI, CSC: CuEq calculated from individual metal production guidance midpoints using live IRESS spot prices "
         "(Copper COPNY.ID, Gold SPTGLD.IF, Silver SPTSLV.IF, Zinc ZINC.LME, Lead LEAD.LME). "
         "Figures will change in real-time as spot prices move."),
        ("§  Pre-production run-rate (CYM)",
         "Cyprium Metals is pre-production. Phase 1 Nifty SX-EW cathode restart is targeting first production in Q3 CY2026 "
         "at a run-rate of approximately 7,000 tpa Cu cathode. CuEq figure shown is the annualised Phase 1 run-rate, not FY2026 actual production. "
         "EV/CuEq for CYM should be interpreted as a development-stage multiple, not a producing-company multiple."),
        ("CSC note",
         "Capstone Copper Corp. is a Canadian company listed on ASX via CHESS Depository Interests (CDIs, 1:1 ratio). "
         "All financials are in USD converted to AUD at 0.645 (approximate implied rate at data sourcing). "
         "Capstone does not file ASX Appendix 5B; financial data is sourced from Canadian MD&A (Q1 CY2026, 29-Apr-2026). "
         "Its scale (200–230 kt Cu guidance) is substantially larger than other peers in this table."),
        ("DVP note",
         "Develop Global's Woodlawn annual copper and zinc guidance figures are LOM run-rate estimates. "
         "Formal FY2026 annual Cu/Zn guidance in kt has not been published as a specific annual target. "
         "Pb and Ag co-product guidance is not quantified and has been excluded from the CuEq calculation."),
        ("Net (Cash)/Debt",
         "Shown using accounting convention: positive = net debt position; parenthetical (negative) = net cash position. "
         "Figures from most recent Appendix 5B quarterly cashflow report (March 2026 quarter). "
         "USD-denominated debt (29M: US$120M RCF; CYM: US$27.3M Nebari; CSC: US$1,131.8M) converted at 0.645 AUDUSD."),
        ("Data pull date",
         "Financial data and production guidance sourced from ASX announcements released April–May 2026 "
         "(reflecting March 2026 quarter). Share prices are live via IRESS RTD. Metal prices are live via IRESS RTD. "
         "This table refreshes automatically when IRESS is connected."),
    ]

    merge_title(ws, fn_row + 1, 1, TOTAL_COLS,
                "  FOOTNOTES & DISCLOSURE NOTES",
                sz=9, bg=C_NAVY, row_ht=14)
    fn_row += 2

    for label, text in footnotes:
        ws.row_dimensions[fn_row].height = 14
        write(ws, fn_row, 1, val=label,
              bold=True, sz=8, col_txt=C_AMBER_DARK, bg=C_AMBER,
              h="left", no_border=True)
        ws.merge_cells(start_row=fn_row, start_column=2,
                       end_row=fn_row,   end_column=TOTAL_COLS)
        write(ws, fn_row, 2, val=text,
              sz=8, italic=True, bg=C_AMBER, col_txt=C_AMBER_DARK,
              h="left", wrap=True, no_border=True)
        fn_row += 1

    # ── Live price strip (bottom banner) ─────────────────────────────────────
    fn_row += 1
    merge_title(ws, fn_row, 1, TOTAL_COLS,
                "  LIVE METAL PRICE MONITOR  (via IRESS RTD)",
                sz=9, bg=C_NAVY2, row_ht=14)
    fn_row += 1
    ws.row_dimensions[fn_row].height = 18
    price_strip = [
        ("Cu (USD/t)",  f'=@IRESSRtd("Quote","13","0","COPNY.ID")',  '"US$"#,##0'),
        ("Au (USD/oz)", f'=@IRESSRtd("Quote","13","0","SPTGLD.IF")', '"US$"#,##0'),
        ("Ag (USD/oz)", f'=@IRESSRtd("Quote","13","0","SPTSLV.IF")', '"US$"#,##0.00'),
        ("Zn (USD/t)",  f'=@IRESSRtd("Quote","13","0","ZINC.LME")',  '"US$"#,##0'),
        ("Pb (USD/t)",  f'=@IRESSRtd("Quote","13","0","LEAD.LME")',  '"US$"#,##0'),
        ("AUD/USD",     f'=@IRESSRtd("Quote","13","0","AUDUSD.FX")', "0.0000"),
    ]
    # pair label/value across 2 cols each (using cols 1-12 for 6 pairs = cols 1-2, 3-4, ... but we have 11 cols)
    # Lay out as: label in odd cols, value in even cols (across ~10 cols for 5 items, AUD/USD in last)
    strip_cols = [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12]
    # Use cols 1-2, 3-4, 5-6, 7-8, 9-10, 11-12 but we only have 11 visible cols
    label_cols = [1, 3, 5, 7, 9, 11]
    val_cols   = [2, 4, 6, 8, 10, 11]
    for j, (lbl, frm, fmt) in enumerate(price_strip):
        lc_ = label_cols[j] if j < 5 else 10
        vc_ = val_cols[j]   if j < 5 else 11
        write(ws, fn_row, lc_, val=lbl,
              bold=True, sz=8, col_txt=C_DGREY, bg=C_BGASMP, h="right")
        write(ws, fn_row, vc_, formula=frm,
              bold=True, sz=9, col_txt=C_NAVY, bg=C_BGASMP, h="left", num_fmt=fmt)

    # ── Freeze panes (header row + 1) ────────────────────────────────────────
    ws.freeze_panes = ws.cell(row=COMPS_DATA_START, column=CC_NAME)

    # ── Print setup ───────────────────────────────────────────────────────────
    ws.page_setup.orientation    = "landscape"
    ws.page_setup.fitToPage      = True
    ws.page_setup.fitToWidth     = 1
    ws.page_setup.fitToHeight    = 0
    ws.print_title_rows          = f"$1:${COMPS_HDR_ROW}"
    ws.sheet_properties.tabColor = C_NAVY


# ──────────────────────────────────────────────────────────────────────────────
# MAIN
# ──────────────────────────────────────────────────────────────────────────────
def main():
    wb = openpyxl.Workbook()

    # Sheet 1 – Comps Table
    ws_comps = wb.active
    ws_comps.title = "Comps Table"
    ws_comps.sheet_properties.tabColor = C_NAVY

    # Sheet 2 – Assumptions
    ws_asmpt = wb.create_sheet("Assumptions")
    ws_asmpt.sheet_properties.tabColor = "2E75B6"

    build_assumptions(ws_asmpt)
    build_comps(ws_comps)

    out = "copper_comps_table.xlsx"
    wb.save(out)
    print(f"✓  Saved: {out}")
    print(f"   {len(COMPANIES)} companies | 2 sheets | IRESS RTD formulas active on open in Excel")


if __name__ == "__main__":
    main()
